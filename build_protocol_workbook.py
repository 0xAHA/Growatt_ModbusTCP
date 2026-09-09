#!/usr/bin/env python3
"""Build a reviewable workbook of protocol registers beside our current mappings.

    python3 build_protocol_workbook.py

Writes `Protocols/protocol_register_review.xlsx`, one sheet per protocol and register
space, with the vendor's own columns on the left and what this integration currently does
on the right, so the two can be compared row by row.

Why this exists
---------------
The markdown extractions under `docs/developer/` are readable but lossy. `protocol-v139.md`
drops the **Range** column, which is the only place V1.39 says a register can go negative -
53 registers use it. So the extracted reference could not answer "is this register signed?",
and that question is exactly what produced #429: register 36 decoded unsigned, every
negative reading discarded as a 429 MW underflow, and the symptom a *missing* value rather
than a wrong one.

VPP 2.03 is better: it carries an explicit `Type` column (`INT16`, `UINT32`, ...), which
settles signedness outright for the registers it covers.

Reading straight from the vendor workbooks avoids trusting either extraction.

The two sources do not describe registers the same way, so the sheets differ:

    V1.39      Register No | Name | Description | Read/Write | Range | Unit |
               Initial value | Note
    VPP 2.03   No. | Name | Read/Write | Type | Unit | Address | Number | Range

Filling it in
-------------
`Spec signed?`, `We use signed`, and `Check` are derived. **`Verdict` and `Reviewer notes`
are deliberately empty** - they are for the human reading the PDF. Once filled in, this
workbook is the settled reference and beats both the markdown and the source.

`Check` is a prompt, not a judgement. `SIGN MISMATCH` means the two disagree and someone
should look; it does not mean we are wrong. A directional register can be unsigned in the
document and still need a sign in practice, and the reverse happens too.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required:  pip install openpyxl")

ROOT = Path(__file__).parent
PROTOCOLS = ROOT / "Protocols"
V139 = PROTOCOLS / "Growatt_WIT-Modbus-RTU-Protocol-II-V1.39.xlsx"
VPP = PROTOCOLS / "GI-BK-E060_GROWATT.VPP.COMMUNICATION.PROTOCOL.OF.INVERTER_V2.03.xlsx"
OUTPUT = PROTOCOLS / "protocol_register_review.xlsx"

# V1.39 keeps one logical table per sheet, split by a heading part way down. Columns are
# sparse and shift by a cell or two between rows, so each field is a RANGE of columns and
# the first non-empty cell in it wins.
V139_INPUT_HEADING_ROW = 863          # the "2.2 Input register" heading
V139_HOLDING_COLS = {
    "register": (0, 2), "name": (3, 6), "description": (7, 18), "rw": (19, 22),
    "range": (23, 28), "unit": (29, 32), "initial": (33, 36), "note": (37, 56),
}
V139_INPUT_COLS = {
    "register": (0, 3), "name": (4, 12), "description": (13, 31),
    "range": (32, 37), "unit": (38, 42), "note": (43, 56),
}
VPP_HOLDING_COLS = {
    "no": 0, "name": 1, "rw": 15, "type": 19, "unit": 26, "address": 31,
    "number": 35, "range": 40,
}
VPP_INPUT_COLS = {
    "no": 0, "name": 1, "rw": 2, "type": 3, "unit": 4, "address": 5,
    "number": 6, "range": 7,
}

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
OURS_FILL = PatternFill("solid", fgColor="2E5C3E")
REVIEW_FILL = PatternFill("solid", fgColor="7B3F00")
MISMATCH_FILL = PatternFill("solid", fgColor="FFE0B2")
UNMAPPED_FILL = PatternFill("solid", fgColor="F2F2F2")


def _text(row, index) -> str:
    if index is None or index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).replace("\n", " ").strip()


def _first(row, span) -> str:
    lo, hi = span
    for i in range(lo, min(hi, len(row) - 1) + 1):
        value = _text(row, i)
        if value:
            return value
    return ""


def _looks_negative(text: str) -> bool:
    """True when a Range cell describes a bound below zero, e.g. `-15~15A`.

    A hyphen is only a minus sign at the start of the cell or after a separator. Growatt
    also use it as a range separator - `0-100` is nought to a hundred, not minus one
    hundred - and treating those as signed produced 74 false mismatches on the holding
    sheet alone, which would have buried the real ones.
    """
    return bool(re.search(r"(?:^|[~,;(\[:=\s])-\s*\d", text.strip()))


def read_v139():
    """(holding, input) rows from the V1.39 workbook, continuation lines folded in."""
    book = openpyxl.load_workbook(V139, read_only=True, data_only=True)
    sheet = book["Table 1"]
    holding, inputs, group = [], [], ""

    for number, raw in enumerate(sheet.iter_rows(values_only=True), 1):
        is_input = number > V139_INPUT_HEADING_ROW
        columns = V139_INPUT_COLS if is_input else V139_HOLDING_COLS
        target = inputs if is_input else holding

        register = _first(raw, columns["register"])
        name = _first(raw, columns["name"])

        # The input table is divided into groups ("First group", ...). Worth carrying:
        # the same address means different things in different groups.
        if is_input and register and not register.isdigit() and "group" in register.lower():
            group = register
            continue
        if not register.isdigit():
            # A wrapped cell belonging to the register above rather than a new one.
            if target and (name or _first(raw, columns["description"])):
                previous = target[-1]
                for field in ("name", "description", "note"):
                    extra = _first(raw, columns.get(field, (99, 99)))
                    if extra:
                        previous[field] = f"{previous[field]} {extra}".strip()
            continue

        entry = {
            "register": int(register),
            "name": name,
            "description": _first(raw, columns["description"]),
            "rw": _first(raw, columns["rw"]) if "rw" in columns else "",
            "range": _first(raw, columns["range"]),
            "unit": _first(raw, columns["unit"]),
            "initial": _first(raw, columns["initial"]) if "initial" in columns else "",
            "note": _first(raw, columns["note"]),
            "group": group if is_input else "",
        }
        target.append(entry)
    return holding, inputs


def read_vpp():
    """(holding, input) rows from the VPP 2.03 workbook."""
    book = openpyxl.load_workbook(VPP, read_only=True, data_only=True)
    out = []
    for sheet_name, columns in (("HOLDING REGISTERS", VPP_HOLDING_COLS),
                                ("input registers", VPP_INPUT_COLS)):
        rows, section = [], ""
        for raw in book[sheet_name].iter_rows(min_row=3, values_only=True):
            address = _text(raw, columns["address"])
            first = _text(raw, columns["no"])
            if not address:
                # Section banners look like "PV Parameter:31010".
                if first and not first.isdigit():
                    section = first
                continue
            if not address.isdigit():
                continue
            rows.append({
                "no": first,
                "name": _text(raw, columns["name"]),
                "rw": _text(raw, columns["rw"]),
                "type": _text(raw, columns["type"]),
                "unit": _text(raw, columns["unit"]),
                "register": int(address),
                "number": _text(raw, columns["number"]),
                "range": _text(raw, columns["range"]),
                "group": section,
            })
        out.append(rows)
    return out[0], out[1]


def load_our_maps():
    """(space, address) -> list of what each profile says about it."""
    sys.path.insert(0, str(ROOT / "custom_components" / "growatt_modbus"))
    import profiles  # noqa: E402  - path set above

    ours: dict[tuple[str, int], list[dict]] = {}
    for profile_name, profile in profiles.REGISTER_MAPS.items():
        for space in ("input_registers", "holding_registers"):
            for address, info in (profile.get(space) or {}).items():
                key = ("input" if space == "input_registers" else "holding", address)
                ours.setdefault(key, []).append({"profile": profile_name, **info})
    return ours


def summarise_ours(entries):
    """Collapse every profile's view of one address into cells for the sheet."""
    if not entries:
        return {"name": "", "scale": "", "signed": "", "combined_scale": "",
                "combined_unit": "", "pair": "", "access": "", "count": 0, "profiles": ""}

    def distinct(field):
        seen = []
        for entry in entries:
            value = entry.get(field)
            if value is not None and value not in seen:
                seen.append(value)
        return ", ".join(str(v) for v in seen)

    signed_flags = {bool(e.get("signed")) for e in entries}
    signed = "yes" if signed_flags == {True} else "no" if signed_flags == {False} else "MIXED"
    return {
        "name": distinct("name") or distinct("alias"),
        "scale": distinct("scale"),
        "signed": signed,
        "combined_scale": distinct("combined_scale"),
        "combined_unit": distinct("combined_unit"),
        "pair": distinct("pair"),
        "access": distinct("access"),
        "count": len(entries),
        "profiles": ", ".join(sorted(e["profile"] for e in entries)),
    }


def _span(row) -> list[int]:
    """Every register a spec row covers. `Number` is 1 unless the document says otherwise."""
    start = row["register"]
    try:
        count = int(str(row.get("number") or 1).strip())
    except ValueError:
        count = 1
    return list(range(start, start + max(count, 1)))


def spec_is_signed(row) -> str:
    """What the document says about sign: 'yes', 'no', or '' when it does not say."""
    kind = (row.get("type") or "").upper()
    if kind.startswith("INT"):
        return "yes"
    if kind.startswith(("UINT", "BOOL")):
        return "no"
    if _looks_negative(row.get("range", "")):
        return "yes"
    return ""


def write_sheet(book, title, spec_columns, rows, ours, space):
    sheet = book.create_sheet(title)
    ours_columns = ["Our name", "Scale", "We use signed", "Combined scale",
                    "Combined unit", "Pair", "Access", "Profiles", "Profile list"]
    review_columns = ["Spec signed?", "Check", "Verdict", "Reviewer notes"]
    headers = [label for label, _ in spec_columns] + ours_columns + review_columns
    sheet.append(headers)

    for index, label in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if index <= len(spec_columns):
            cell.fill = HEADER_FILL
        elif index <= len(spec_columns) + len(ours_columns):
            cell.fill = OURS_FILL
        else:
            cell.fill = REVIEW_FILL

    for row in rows:
        # A spec row can cover several registers (VPP's `Number`), and for a 32-bit pair
        # we carry `signed` on the LOW word because that is where combined_scale lives.
        # Comparing only the first address called every such pair a mismatch.
        entries = []
        for address in _span(row):
            entries.extend(ours.get((space, address), []))
        mine = summarise_ours(entries)
        declared = spec_is_signed(row)

        if not entries:
            check = "not mapped"
        elif declared == "yes" and mine["signed"] == "no":
            check = "SIGN MISMATCH - spec allows negative, we decode unsigned"
        elif declared == "no" and mine["signed"] == "yes":
            check = "SIGN MISMATCH - spec is unsigned, we decode signed"
        elif mine["signed"] == "MIXED":
            check = "profiles disagree with each other"
        elif declared == "":
            check = "spec does not say"
        else:
            check = "agrees"

        sheet.append(
            [row.get(key, "") for _, key in spec_columns]
            + [mine["name"], mine["scale"], mine["signed"], mine["combined_scale"],
               mine["combined_unit"], mine["pair"], mine["access"],
               mine["count"] or "", mine["profiles"]]
            + [declared, check, "", ""]
        )

        written = sheet.max_row
        if check.startswith("SIGN MISMATCH") or check.startswith("profiles disagree"):
            for column in range(1, len(headers) + 1):
                sheet.cell(row=written, column=column).fill = MISMATCH_FILL
        elif not entries:
            for column in range(1, len(headers) + 1):
                sheet.cell(row=written, column=column).fill = UNMAPPED_FILL

    widths = [10, 26, 40, 8, 20, 10, 12, 34] + [24, 10, 13, 14, 14, 8, 9, 9, 40] + [12, 46, 14, 40]
    for index in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = (
            widths[index - 1] if index - 1 < len(widths) else 18)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"
    return sheet


def write_extras(book, ours, covered):
    """Addresses we map that the documents do not list, so nothing is invisible."""
    sheet = book.create_sheet("Ours not in spec")
    headers = ["Space", "Register", "Our name", "Scale", "We use signed",
               "Combined scale", "Combined unit", "Pair", "Access", "Profiles",
               "Profile list", "Verdict", "Reviewer notes"]
    sheet.append(headers)
    for index in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = OURS_FILL if index <= 11 else REVIEW_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for (space, address) in sorted(set(ours) - covered):
        mine = summarise_ours(ours[(space, address)])
        sheet.append([space, address, mine["name"], mine["scale"], mine["signed"],
                      mine["combined_scale"], mine["combined_unit"], mine["pair"],
                      mine["access"], mine["count"], mine["profiles"], "", ""])

    for index, width in enumerate([9, 10, 26, 10, 13, 14, 14, 8, 9, 9, 40, 14, 40], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"
    return sheet


def write_readme(book, counts):
    sheet = book.create_sheet("README", 0)
    lines = [
        ("Growatt protocol register review", True),
        ("", False),
        ("Vendor columns on the left, what this integration currently does in the middle,", False),
        ("and two empty columns on the right for your conclusions.", False),
        ("", False),
        ("Generated by build_protocol_workbook.py, read straight from the workbooks in", False),
        ("Protocols/ rather than from the markdown under docs/developer/ - the V1.39", False),
        ("extraction drops the Range column, which is the only place that document says a", False),
        ("register may go negative.", False),
        ("", False),
        ("Sheets", True),
        ("  V1.39 Holding / V1.39 Input      Register No, Name, Description, Read/Write,", False),
        ("                                   Range, Unit, Initial value, Note", False),
        ("  VPP2.03 Holding / VPP2.03 Input  No., Name, Read/Write, Type, Unit, Address,", False),
        ("                                   Number, Range", False),
        ("  Ours not in spec                 addresses we map that neither document lists", False),
        ("", False),
        ("Columns to fill in", True),
        ("  Verdict          your conclusion, e.g. signed / unsigned / leave as is", False),
        ("  Reviewer notes   anything worth keeping with it", False),
        ("", False),
        ("Derived columns", True),
        ("  Spec signed?     yes when the Type is INT*, or the Range has a negative bound.", False),
        ("                   Blank means the document does not say - which is the case for", False),
        ("                   most of V1.39, including registers 35/36 and 3041-3044.", False),
        ("  We use signed    from the profiles. MIXED means profiles disagree.", False),
        ("  Check            a prompt, not a judgement. SIGN MISMATCH means the two", False),
        ("                   disagree and someone should look - a directional register can", False),
        ("                   be unsigned on paper and still need a sign in practice.", False),
        ("", False),
        ("Colours", True),
        ("  Amber   the two disagree, or our own profiles disagree with each other", False),
        ("  Grey    the document lists it and we do not map it", False),
        ("", False),
        ("Counts at generation", True),
    ]
    for label, value in counts:
        lines.append((f"  {label}: {value}", False))
    lines += [
        ("", False),
        ("Note that V1.39 states, in section 2: \"It is 16 bits (two bytes) unsigned", False),
        ("integer for each holding and input register.\" Unsigned is the documented", False),
        ("default, so every signed flag we carry is a departure from it that needs its own", False),
        ("justification.", False),
    ]
    for label, bold in lines:
        sheet.append([label])
        if bold:
            sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True, size=12)
    sheet.column_dimensions["A"].width = 100
    return sheet


def main() -> int:
    for path in (V139, VPP):
        if not path.exists():
            sys.exit(f"missing source workbook: {path}")

    v139_holding, v139_input = read_v139()
    vpp_holding, vpp_input = read_vpp()
    ours = load_our_maps()

    v139_spec = [("Register No", "register"), ("Name", "name"), ("Description", "description"),
                 ("Read/Write", "rw"), ("Range", "range"), ("Unit", "unit"),
                 ("Initial value", "initial"), ("Note", "note")]
    v139_input_spec = v139_spec + [("Group", "group")]
    vpp_spec = [("No.", "no"), ("Name", "name"), ("Read/Write", "rw"), ("Type", "type"),
                ("Unit", "unit"), ("Address", "register"), ("Number", "number"),
                ("Range", "range"), ("Section", "group")]

    book = openpyxl.Workbook()
    book.remove(book.active)

    covered = set()
    for title, spec, rows, space in (
        ("V1.39 Holding", v139_spec, v139_holding, "holding"),
        ("V1.39 Input", v139_input_spec, v139_input, "input"),
        ("VPP2.03 Holding", vpp_spec, vpp_holding, "holding"),
        ("VPP2.03 Input", vpp_spec, vpp_input, "input"),
    ):
        write_sheet(book, title, spec, rows, ours, space)
        covered |= {(space, address) for row in rows for address in _span(row)}

    write_extras(book, ours, covered)
    write_readme(book, [
        ("V1.39 holding rows", len(v139_holding)),
        ("V1.39 input rows", len(v139_input)),
        ("VPP2.03 holding rows", len(vpp_holding)),
        ("VPP2.03 input rows", len(vpp_input)),
        ("addresses we map", len(ours)),
        ("addresses we map that neither document lists", len(set(ours) - covered)),
    ])

    book.save(OUTPUT)
    print(f"wrote {OUTPUT.relative_to(ROOT)}")
    for sheet in book.sheetnames:
        print(f"  {sheet}: {book[sheet].max_row - 1} row(s)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
